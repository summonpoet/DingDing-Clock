"""
报告生成模块
支持生成 Markdown 报告和 Excel 表格
"""

import os
from datetime import datetime

from .scraper import InvestEvent
from .researcher import CompanyResearch


def generate_markdown_report(
    results: list[CompanyResearch],
    output_dir: str = "output",
) -> str:
    """
    生成 Markdown 格式的投研报告

    Args:
        results: 公司研究结果列表
        output_dir: 输出目录

    Returns:
        生成的文件路径
    """
    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(output_dir, f"投研报告_{timestamp}.md")

    lines = []
    lines.append(f"# 投资事件研究报告")
    lines.append(f"")
    lines.append(f"> 生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"> 共分析 {len(results)} 家公司")
    lines.append(f"")

    # 目录
    lines.append("## 目录\n")
    for i, r in enumerate(results, 1):
        event = r.event
        lines.append(f"{i}. [{event.company}](#{_anchor(event.company)}) - {event.industry} | {event.round} | {event.amount}")
    lines.append("")

    # 总览表格
    lines.append("## 投资事件总览\n")
    lines.append("| # | 日期 | 公司 | 行业 | 轮次 | 金额 | 投资方 | 估值 |")
    lines.append("|---|------|------|------|------|------|--------|------|")
    for i, r in enumerate(results, 1):
        e = r.event
        lines.append(f"| {i} | {e.date} | {e.company} | {e.industry} | {e.round} | {e.amount} | {e.investors} | {e.valuation} |")
    lines.append("")

    # 分隔线
    lines.append("---\n")

    # 每家公司的详细分析
    for i, r in enumerate(results, 1):
        event = r.event
        lines.append(f"## {i}. {event.company} {{#{_anchor(event.company)}}}\n")

        # 基本信息卡片
        lines.append("### 融资信息\n")
        lines.append(f"| 字段 | 信息 |")
        lines.append(f"|------|------|")
        lines.append(f"| 行业 | {event.industry} |")
        lines.append(f"| 轮次 | {event.round} |")
        lines.append(f"| 金额 | {event.amount} |")
        lines.append(f"| 投资方 | {event.investors} |")
        lines.append(f"| 估值 | {event.valuation} |")
        lines.append(f"| 日期 | {event.date} |")
        lines.append("")

        # AI 分析内容
        lines.append("### AI 研究分析\n")
        lines.append(r.raw_response if r.raw_response else r.summary)
        lines.append("")

        # 分隔线
        if i < len(results):
            lines.append("\n---\n")

    content = "\n".join(lines)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

    return filepath


def generate_excel_report(
    results: list[CompanyResearch],
    output_dir: str = "output",
) -> str:
    """
    生成 Excel 格式的投研报告

    Args:
        results: 公司研究结果列表
        output_dir: 输出目录

    Returns:
        生成的文件路径
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(output_dir, f"投研报告_{timestamp}.xlsx")

    wb = Workbook()

    # ===== Sheet 1: 投资事件总览 =====
    ws1 = wb.active
    ws1.title = "投资事件总览"

    # 样式定义
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 表头
    headers = ["序号", "日期", "公司名称", "行业", "融资轮次", "融资金额", "投资方", "最新估值"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 数据行
    for row, r in enumerate(results, 2):
        e = r.event
        values = [row - 1, e.date, e.company, e.industry, e.round, e.amount, e.investors, e.valuation]
        for col, val in enumerate(values, 1):
            cell = ws1.cell(row=row, column=col, value=val)
            cell.alignment = cell_alignment
            cell.border = thin_border

    # 列宽
    col_widths = [6, 12, 20, 15, 12, 15, 30, 15]
    for i, width in enumerate(col_widths, 1):
        ws1.column_dimensions[chr(64 + i)].width = width

    # ===== Sheet 2: 详细研究报告 =====
    ws2 = wb.create_sheet("详细研究报告")

    headers2 = ["序号", "公司名称", "行业", "融资轮次", "融资金额", "AI研究摘要"]
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row, r in enumerate(results, 2):
        e = r.event
        summary = r.raw_response if r.raw_response else r.summary
        # Excel 单元格有字符限制，截断过长的内容
        if len(summary) > 32000:
            summary = summary[:32000] + "\n...(内容过长已截断)"

        values = [row - 1, e.company, e.industry, e.round, e.amount, summary]
        for col, val in enumerate(values, 1):
            cell = ws2.cell(row=row, column=col, value=val)
            cell.alignment = cell_alignment
            cell.border = thin_border

    col_widths2 = [6, 20, 15, 12, 15, 80]
    for i, width in enumerate(col_widths2, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(64 + (i - 1) % 26 + 1)
        ws2.column_dimensions[col_letter].width = width

    wb.save(filepath)
    return filepath


def _anchor(text: str) -> str:
    """生成Markdown锚点ID"""
    return text.lower().replace(" ", "-").replace("(", "").replace(")", "")
